"""Bounded, dependency-light readers for library CSV/TSV/ODS/XLSX exports."""

from __future__ import annotations

import csv
from pathlib import Path
import zipfile

from defusedxml import ElementTree

MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_XML_BYTES = 32 * 1024 * 1024
MAX_ROWS = 20001
MAX_COLUMNS = 128
MAX_CELL_CHARS = 4096
TABLE_NS = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
OFFICE_NS = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"


def _bounded(row: list) -> list:
    return [str(value or "")[:MAX_CELL_CHARS] for value in row[:MAX_COLUMNS]]


def read_table(path: str | Path) -> list[list]:
    """Read a supported export; reject oversize, malformed and unsupported input."""
    path = Path(path)
    if not path.is_file():
        raise ValueError("Spreadsheet file does not exist.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Spreadsheet exceeds the 16 MiB input limit.")
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = []
            for row in csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ","):
                if len(rows) >= MAX_ROWS:
                    raise ValueError("Spreadsheet exceeds the row limit.")
                rows.append(_bounded(row))
            return rows
    if suffix not in {".ods", ".xlsx"}:
        raise ValueError("Use CSV, TSV, ODS or XLSX.")
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XML_BYTES:
                raise ValueError("Expanded spreadsheet exceeds the 32 MiB limit.")
            if suffix == ".ods":
                xml = archive.read("content.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("Spreadsheet is invalid or is an unhydrated Git LFS pointer.") from exc
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            rows = []
            for row in book.active.iter_rows(values_only=True):
                if len(rows) >= MAX_ROWS:
                    raise ValueError("Spreadsheet exceeds the row limit.")
                rows.append(_bounded(list(row)))
            return rows
        finally:
            book.close()
    root = ElementTree.fromstring(xml)
    table = root.find(f".//{{{TABLE_NS}}}table")
    if table is None:
        raise ValueError("ODS contains no sheet.")
    rows = []
    for element in table.iter(f"{{{TABLE_NS}}}table-row"):
        row = []
        for cell in element:
            if cell.tag not in {f"{{{TABLE_NS}}}table-cell", f"{{{TABLE_NS}}}covered-table-cell"}:
                continue
            text = " ".join("".join(p.itertext()) for p in cell.findall(f".//{{{TEXT_NS}}}p"))
            text = text or cell.get(f"{{{OFFICE_NS}}}value", "")
            repeat = min(int(cell.get(f"{{{TABLE_NS}}}number-columns-repeated", "1")), MAX_COLUMNS)
            row.extend([text[:MAX_CELL_CHARS]] * min(repeat, MAX_COLUMNS - len(row)))
        while row and not row[-1]:
            row.pop()
        if not row:
            continue
        repeat = int(element.get(f"{{{TABLE_NS}}}number-rows-repeated", "1"))
        if repeat < 1 or len(rows) + repeat > MAX_ROWS:
            raise ValueError("ODS exceeds the row limit.")
        rows.extend([list(row) for _ in range(repeat)])
    return rows
